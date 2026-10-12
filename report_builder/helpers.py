# -*- coding: utf-8 -*-
"""
Helper functions for formatting the academic report according to university standards:
- Font: Times New Roman, Unicode, Size 14pt (Pt(14))
- Margins: Left 3cm, Right 2cm, Top 2cm, Bottom 2cm
- Line spacing: Multiple 1.3
- Text alignment: Justified (Both sides)
- First line indent: 1.27 cm
- Chapter Title: Bold, UPPERCASE, Size 15pt, Centered, starts on new page
- Heading 2: Bold, Size 14pt, Left aligned
- Heading 3: Bold, Italic, Size 14pt, Left aligned
- Heading 4: Italic, Size 14pt, Left aligned
- Table Caption: ABOVE table, Centered, Bold
- Figure Caption: BELOW figure, Centered, Bold Italic
- Page number: In footer center, starts from MỞ ĐẦU
"""

import os
import docx
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def apply_section_margins(section):
    """Set Left 3cm, Right 2cm, Top 2cm, Bottom 2cm."""
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)

def set_page_number_start(section, start_num=1):
    """Set section page number starting value."""
    sectPr = section._sectPr
    pgNumType = parse_xml(r'<w:pgNumType %s w:start="%d"/>' % (nsdecls('w'), start_num))
    sectPr.append(pgNumType)

def add_page_number_to_footer(footer):
    """Add centered dynamic PAGE field to footer."""
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(50, 50, 50)

def add_chapter_title(doc, chapter_num_str, title_str):
    """Chapter Title: Bold, UPPERCASE, Size 15, Centered."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(f"CHƯƠNG {chapter_num_str}. {title_str.upper()}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_special_title(doc, title_str):
    """Special Title like MỞ ĐẦU, KẾT LUẬN, TÀI LIỆU THAM KHẢO, PHỤ LỤC."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(title_str.upper())
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    """Heading level 2: Bold, Size 14, Left aligned."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, text):
    """Heading level 3: Bold, Italic, Size 14, Left aligned."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_4(doc, text):
    """Heading level 4: Italic, Size 14, Left aligned."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

import re

def add_body_p(doc, text, is_bullet=False):
    """Body paragraph: Justified, Line spacing 1.3, Indent 1.27cm, Size 14."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    
    # Check if text starts with bullet-like chars (•, *, +, -, –)
    bullet_pattern = r'^\s*([•\*\+\-–])\s*'
    has_bullet = bool(re.match(bullet_pattern, text))
    if is_bullet or has_bullet:
        text = re.sub(bullet_pattern, '', text)
        p.paragraph_format.left_indent = Cm(1.27)
        p.paragraph_format.first_line_indent = Cm(-0.63)
        run_bullet = p.add_run("– ")
        run_bullet.font.name = 'Times New Roman'
        run_bullet.font.size = Pt(14)
        run_bullet.font.color.rgb = RGBColor(0, 0, 0)
    else:
        p.paragraph_format.first_line_indent = Cm(1.27)
        
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_table_caption(doc, table_id_str, title_str):
    """Table Caption ABOVE table, Centered, Bold, Size 13."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(f"Bảng {table_id_str}. {title_str}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_figure_caption(doc, figure_id_str, title_str):
    """Figure Caption BELOW figure, Centered, Bold Italic, Size 13."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(f"Hình {figure_id_str}. {title_str}")
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.bold = True
    run.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_image_box(doc, image_path, width_inches=6.0):
    """Insert centered picture."""
    if os.path.exists(image_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        run = p.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        return p
    return None

def set_cell_background(cell, fill_hex="FFFFFF"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def format_table_row(row, col_widths, bg_color=None, is_header=False, font_size=12, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Format row: default white background, pure black text (RGB 0,0,0)."""
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        # Always use white background as requested
        set_cell_background(cell, "FFFFFF")
        for p in cell.paragraphs:
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Cm(0)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(font_size)
                if is_header:
                    r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)

def add_code_block(doc, code_str, caption=None):
    """
    Insert code snippet with proper formatting:
    - Pure white background, single black/dark border
    - Font Consolas 10pt, pure black text (RGB 0,0,0)
    - Left-aligned (WD_ALIGN_PARAGRAPH.LEFT) to prevent justified whitespace stretching
    - Line spacing 1.15
    """
    code_clean = code_str.strip()
    if code_clean.startswith("```"):
        lines = code_clean.split("\n")
        if lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        code_clean = "\n".join(lines)
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.3)
    set_cell_background(cell, "FFFFFF")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    set_table_borders(tbl, color="333333", sz="4")
    
    p_first = cell.paragraphs[0]
    code_lines = code_clean.split("\n")
    for idx, line in enumerate(code_lines):
        p = p_first if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        
        run = p.add_run(line if line else " ")
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(8)
        p_cap.paragraph_format.first_line_indent = Cm(0)
        r_c = p_cap.add_run(caption)
        r_c.font.name = 'Times New Roman'
        r_c.font.size = Pt(11)
        r_c.italic = True
        r_c.font.color.rgb = RGBColor(0, 0, 0)

